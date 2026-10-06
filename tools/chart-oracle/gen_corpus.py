"""Generate the chart oracle corpus: one PPTX (one chart per slide) and one
XLSX (one chart per sheet) whose chart parts are written by hand, so every
case states exactly the XML Office is asked to draw. Office's PDF export of
these files is the ground truth the engine is measured against.

Two flavours of XML:
  office — what PowerPoint/Excel 365 write when you insert a chart
           (explicit spPr/txPr everywhere, c:style 2 + chartStyle parts omitted)
  bare   — the minimum the schema allows, so Office must fall back to its own
           defaults (c:style, palette, 10pt text). Third-party generators write this.
"""
import json, os, sys, zipfile
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from lxml import etree

NS = ('xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')

GRAY_TX = '<a:solidFill><a:schemeClr val="tx1"><a:lumMod val="65000"/><a:lumOff val="35000"/></a:schemeClr></a:solidFill>'
GRID_LN = '<a:ln w="9525" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="tx1"><a:lumMod val="15000"/><a:lumOff val="85000"/></a:schemeClr></a:solidFill><a:round/></a:ln>'


def txpr(sz=None, bold=False, rot=None, color=GRAY_TX, flavour="office"):
    if flavour == "bare" and sz is None and rot is None:
        return ""
    attrs = []
    if sz is not None:
        attrs.append(f'sz="{sz}"')
    if bold:
        attrs.append('b="1"')
    body = '<a:bodyPr' + (f' rot="{rot}" vert="horz"' if rot is not None else '') + '/>'
    fill = color if flavour == "office" else ""
    return (f'<c:txPr>{body}<a:lstStyle/><a:p><a:pPr><a:defRPr {" ".join(attrs)}>{fill}'
            f'<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/></a:defRPr></a:pPr>'
            f'<a:endParaRPr lang="en-US"/></a:p></c:txPr>') if flavour == "office" else (
            f'<c:txPr>{body}<a:lstStyle/><a:p><a:pPr><a:defRPr {" ".join(attrs)}/></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>')


def str_cache(values, ref):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{escape(str(v))}</c:v></c:pt>' for i, v in enumerate(values) if v is not None)
    return f'<c:strRef><c:f>{ref}</c:f><c:strCache><c:ptCount val="{len(values)}"/>{pts}</c:strCache></c:strRef>'


def num_cache(values, ref, fmt="General"):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(values) if v is not None)
    return f'<c:numRef><c:f>{ref}</c:f><c:numCache><c:formatCode>{escape(fmt)}</c:formatCode><c:ptCount val="{len(values)}"/>{pts}</c:numCache></c:numRef>'


def col_letter(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


ACCENTS = ["accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]


def series_sppr(kind, i, case, flavour):
    if flavour == "bare":
        return ""
    color = case.get("colors", {}).get(i)
    clr = f'<a:srgbClr val="{color}"/>' if color else f'<a:schemeClr val="{ACCENTS[i % 6]}"/>'
    if kind in ("line", "scatter", "radar") and not case.get("area"):
        if "lineW" in case and case["lineW"] is None:
            return '<c:spPr><a:ln w="25400" cap="rnd"><a:noFill/><a:round/></a:ln><a:effectLst/></c:spPr>'
        w = case.get("lineW", 28575)
        return f'<c:spPr><a:ln w="{w}" cap="rnd"><a:solidFill>{clr}</a:solidFill><a:round/></a:ln><a:effectLst/></c:spPr>'
    if kind == "pie":
        return ""
    return f'<c:spPr><a:solidFill>{clr}</a:solidFill><a:ln><a:noFill/></a:ln><a:effectLst/></c:spPr>'


def marker_xml(kind, i, case, flavour):
    if kind not in ("line", "scatter", "radar"):
        return ""
    m = case.get("marker")
    if m is None:
        return '<c:marker><c:symbol val="none"/></c:marker>' if flavour == "office" else ""
    sym, size = m
    if flavour == "bare":
        return f'<c:marker><c:symbol val="{sym}"/><c:size val="{size}"/></c:marker>'
    clr = f'<a:schemeClr val="{ACCENTS[i % 6]}"/>'
    return (f'<c:marker><c:symbol val="{sym}"/><c:size val="{size}"/><c:spPr><a:solidFill>{clr}</a:solidFill>'
            f'<a:ln w="9525"><a:solidFill>{clr}</a:solidFill></a:ln><a:effectLst/></c:spPr></c:marker>')


def dlbls_xml(case, flavour, kind):
    d = case.get("dLbls")
    if not d:
        return ('<c:dLbls><c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/>'
                '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>') if flavour == "office" else ""
    pos = f'<c:dLblPos val="{d["pos"]}"/>' if d.get("pos") else ""
    numfmt = f'<c:numFmt formatCode="{escape(d["fmt"])}" sourceLinked="0"/>' if d.get("fmt") else ""
    tx = txpr(d.get("sz", 1197 if flavour == "office" else None), flavour=flavour)
    sp = '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln><a:effectLst/></c:spPr>' if flavour == "office" else ""
    return (f'<c:dLbls>{numfmt}{sp}{tx}{pos}<c:showLegendKey val="0"/><c:showVal val="{int(d.get("val", True))}"/>'
            f'<c:showCatName val="{int(d.get("cat", False))}"/><c:showSerName val="0"/>'
            f'<c:showPercent val="{int(d.get("pct", False))}"/><c:showBubbleSize val="0"/>'
            + ('<c:showLeaderLines val="1"/>' if kind == "pie" else '') + '</c:dLbls>')


def series_xml(kind, case, flavour, idx_offset=0, series=None):
    cats = case["cats"]
    # The sheet the cell references point at; `write_case_data` fills the same cells.
    sheet = case.get("sheetRef", "Sheet1")
    out = []
    for i, s in enumerate(series if series is not None else case["series"]):
        si = i + idx_offset
        col = col_letter(si + 1)
        name = f'<c:tx>{str_cache([s["name"]], f"{sheet}!${col}$1")}</c:tx>'
        sp = series_sppr(kind, si, case, flavour)
        extra = ""
        if kind == "bar":
            extra = '<c:invertIfNegative val="0"/>'
        mk = marker_xml(kind, si, case, flavour)
        dpts = ""
        if kind == "pie" and flavour == "office":
            for p in range(len(cats)):
                ex = f'<c:explosion val="{case["explode"][p]}"/>' if case.get("explode", {}).get(p) else ""
                dpts += (f'<c:dPt><c:idx val="{p}"/><c:bubble3D val="0"/>{ex}<c:spPr><a:solidFill><a:schemeClr val="{ACCENTS[p % 6]}"/></a:solidFill>'
                         f'<a:ln w="19050"><a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:ln><a:effectLst/></c:spPr></c:dPt>')
        n = len(cats) + 1
        # Raw XML a case adds to each series: `serExtra` after the points (dLbls, trendline,
        # errBars, in schema order), `serExt` as the series' extLst. Each is a string, or a
        # list with one string per series.
        pick = lambda v: (v[si] if si < len(v) else "") if isinstance(v, list) else (v or "")
        dpts += pick(case.get("serExtra"))
        tail_ext = pick(case.get("serExt"))
        if kind in ("scatter", "bubble"):
            data = (f'<c:xVal>{num_cache(case["xs"], f"{sheet}!$A$2:$A${n}")}</c:xVal>'
                    f'<c:yVal>{num_cache(s["values"], f"{sheet}!${col}$2:${col}${n}", case.get("srcFmt", "General"))}</c:yVal>')
            if kind == "bubble":
                data += f'<c:bubbleSize>{num_cache(s["sizes"], f"{sheet}!$Z$2:$Z${n}")}</c:bubbleSize><c:bubble3D val="0"/>'
            smooth = f'<c:smooth val="{int(case.get("smooth", False))}"/>' if kind == "scatter" else ""
            out.append(f'<c:ser><c:idx val="{si}"/><c:order val="{si}"/>{name}{sp}{mk}{dpts}{data}{smooth}{tail_ext}</c:ser>')
            continue
        # Categories: text by default; `catNums` (with `catFmt`) for dates and numbers; `catXml` replaces the element (multi-level).
        if case.get("catXml"):
            cat = case["catXml"]
        elif case.get("catNums"):
            cat = f'<c:cat>{num_cache(case["catNums"], f"{sheet}!$A$2:$A${n}", case.get("catFmt", "m/d/yyyy"))}</c:cat>'
        else:
            cat = f'<c:cat>{str_cache(cats, f"{sheet}!$A$2:$A${n}")}</c:cat>'
        val = f'<c:val>{num_cache(s["values"], f"{sheet}!${col}$2:${col}${n}", case.get("srcFmt", "General"))}</c:val>'
        smooth = f'<c:smooth val="{int(case.get("smooth", False))}"/>' if kind in ("line", "stock") else ""
        out.append(f'<c:ser><c:idx val="{si}"/><c:order val="{si}"/>{name}{sp}{extra}{mk}{dpts}{cat}{val}{smooth}{tail_ext}</c:ser>')
    return "".join(out)


def axis_xml(tag, ax_id, cross_id, pos, case, flavour, which, deleted=False, cross_between="between"):
    a = case.get(which, {})
    # A category axis may be a date axis (`"tag": "c:dateAx"`, with `baseTimeUnit`, `majorUnit`, `majorTimeUnit`).
    tag = a.get("tag", tag)
    scaling = f'<c:orientation val="{a.get("orientation", "minMax")}"/>'
    if a.get("logBase"):
        scaling = f'<c:logBase val="{a["logBase"]}"/>' + scaling
    if a.get("max") is not None:
        scaling += f'<c:max val="{a["max"]}"/>'
    if a.get("min") is not None:
        scaling += f'<c:min val="{a["min"]}"/>'
    grid = ""
    if a.get("major", which == "valAx" and flavour == "office"):
        grid += f'<c:majorGridlines>{"<c:spPr>" + GRID_LN + "<a:effectLst/></c:spPr>" if flavour == "office" else ""}</c:majorGridlines>'
    if a.get("minor"):
        grid += '<c:minorGridlines>' + ('<c:spPr><a:ln w="9525"><a:solidFill><a:schemeClr val="tx1"><a:lumMod val="5000"/><a:lumOff val="95000"/></a:schemeClr></a:solidFill></a:ln></c:spPr>' if flavour == "office" else '') + '</c:minorGridlines>'
    title = ""
    if a.get("title"):
        rot = ' rot="-5400000" vert="horz"' if pos in ("l", "r") else ""
        tsz = ' sz="1330" b="0"' if flavour == "office" else ''
        title = (f'<c:title><c:tx><c:rich><a:bodyPr{rot}/><a:lstStyle/><a:p><a:pPr><a:defRPr{tsz}/></a:pPr><a:r><a:rPr lang="en-US"/>'
                 f'<a:t>{escape(a["title"])}</a:t></a:r></a:p></c:rich></c:tx><c:overlay val="0"/></c:title>')
    fmt = a.get("fmt")
    numfmt = f'<c:numFmt formatCode="{escape(fmt)}" sourceLinked="0"/>' if fmt else ('<c:numFmt formatCode="General" sourceLinked="1"/>' if flavour == "office" else "")
    major_tick = a.get("majorTick", "none" if flavour == "office" else None)
    ticks = (f'<c:majorTickMark val="{major_tick}"/>' if major_tick else "") + (
        '<c:minorTickMark val="none"/>' if flavour == "office" else "")
    lblpos = f'<c:tickLblPos val="{a.get("tickLblPos", "nextTo")}"/>'
    if flavour == "office":
        line = GRID_LN if (tag != "c:valAx" or a.get("line")) else '<a:ln><a:noFill/></a:ln>'
        sppr = f'<c:spPr><a:noFill/>{line}<a:effectLst/></c:spPr>'
    else:
        sppr = ""
    tx = txpr(a.get("sz", 1197 if flavour == "office" else None), rot=a.get("rot"), flavour=flavour)
    crosses = f'<c:crosses val="{a.get("crosses", "autoZero")}"/>' if a.get("crossesAt") is None else f'<c:crossesAt val="{a["crossesAt"]}"/>'
    tail = ""
    if tag == "c:dateAx":
        tail = '<c:auto val="1"/><c:lblOffset val="100"/>'
        if a.get("baseTimeUnit"):
            tail += f'<c:baseTimeUnit val="{a["baseTimeUnit"]}"/>'
        if a.get("majorUnit"):
            tail += f'<c:majorUnit val="{a["majorUnit"]}"/><c:majorTimeUnit val="{a.get("majorTimeUnit", "days")}"/>'
    elif tag == "c:catAx":
        tail = '<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/>'
        if a.get("skip"):
            tail += f'<c:tickLblSkip val="{a["skip"]}"/>'
        tail += '<c:noMultiLvlLbl val="0"/>'
    elif tag == "c:valAx":
        tail = f'<c:crossBetween val="{cross_between}"/>'
        if a.get("majorUnit"):
            tail += f'<c:majorUnit val="{a["majorUnit"]}"/>'
        if a.get("dispUnits"):
            tail += f'<c:dispUnits><c:builtInUnit val="{a["dispUnits"]}"/><c:dispUnitsLbl/></c:dispUnits>'
    return (f'<{tag}><c:axId val="{ax_id}"/><c:scaling>{scaling}</c:scaling><c:delete val="{int(deleted or a.get("delete", False))}"/>'
            f'<c:axPos val="{pos}"/>{grid}{title}{numfmt}{ticks}{lblpos}{sppr}{tx}<c:crossAx val="{cross_id}"/>{crosses}{tail}</{tag}>')


def chart_xml(case):
    flavour = case.get("flavour", "office")
    kind = case["kind"]
    group = ""
    axes = ""
    plot_layout = "<c:layout/>"
    if case.get("plotLayout"):
        x, y, w, h = case["plotLayout"]
        plot_layout = (f'<c:layout><c:manualLayout><c:layoutTarget val="{case.get("layoutTarget", "inner")}"/><c:xMode val="edge"/><c:yMode val="edge"/>'
                       f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>')
    vary = f'<c:varyColors val="{int(case.get("vary", kind in ("pie", "doughnut")))}"/>'
    dl = dlbls_xml(case, flavour, "pie" if kind in ("pie", "doughnut") else kind)
    if kind == "bar":
        grouping = case.get("grouping", "clustered")
        gap = case.get("gap", 219 if flavour == "office" else None)
        overlap = case.get("overlap", (-27 if grouping == "clustered" else 100) if flavour == "office" else (100 if grouping != "clustered" else None))
        group = (f'<c:barChart><c:barDir val="{case.get("dir", "col")}"/><c:grouping val="{grouping}"/>{vary}'
                 f'{series_xml("bar", case, flavour)}{dl}' + (f'<c:gapWidth val="{gap}"/>' if gap is not None else "")
                 + (f'<c:overlap val="{overlap}"/>' if overlap is not None else "") + '<c:axId val="10"/><c:axId val="20"/></c:barChart>')
        cat_pos, val_pos = ("b", "l") if case.get("dir", "col") == "col" else ("l", "b")
        axes = axis_xml("c:catAx", 10, 20, cat_pos, case, flavour, "catAx") + axis_xml("c:valAx", 20, 10, val_pos, case, flavour, "valAx")
    elif kind in ("line", "area"):
        tag = "lineChart" if kind == "line" else "areaChart"
        grouping = case.get("grouping", "standard")
        mk = '<c:marker val="1"/>' if kind == "line" else ""
        group = (f'<c:{tag}><c:grouping val="{grouping}"/>{vary}{series_xml(kind, case, flavour)}{dl}{mk}'
                 f'<c:axId val="10"/><c:axId val="20"/></c:{tag}>')
        axes = axis_xml("c:catAx", 10, 20, "b", case, flavour, "catAx") + axis_xml(
            "c:valAx", 20, 10, "l", case, flavour, "valAx", cross_between="midCat" if kind == "area" else "between")
    elif kind in ("pie", "doughnut"):
        first = f'<c:firstSliceAng val="{case.get("firstSliceAng", 0)}"/>'
        hole = f'<c:holeSize val="{case.get("hole", 75)}"/>' if kind == "doughnut" else ""
        group = f'<c:{kind}Chart>{vary}{series_xml("pie", case, flavour)}{dl}{first}{hole}</c:{kind}Chart>'
    elif kind == "scatter":
        style = case.get("scatterStyle", "lineMarker")
        group = (f'<c:scatterChart><c:scatterStyle val="{style}"/>{vary}{series_xml("scatter", case, flavour)}{dl}'
                 f'<c:axId val="10"/><c:axId val="20"/></c:scatterChart>')
        axes = axis_xml("c:valAx", 10, 20, "b", case, flavour, "catAx", cross_between="midCat") + axis_xml(
            "c:valAx", 20, 10, "l", case, flavour, "valAx", cross_between="midCat")
    elif kind == "bubble":
        group = (f'<c:bubbleChart>{vary}{series_xml("bubble", case, flavour)}{dl}<c:bubbleScale val="100"/><c:showNegBubbles val="0"/>'
                 f'<c:axId val="10"/><c:axId val="20"/></c:bubbleChart>')
        axes = axis_xml("c:valAx", 10, 20, "b", case, flavour, "catAx", cross_between="midCat") + axis_xml(
            "c:valAx", 20, 10, "l", case, flavour, "valAx", cross_between="midCat")
    elif kind == "stock":
        # Open, high, low, close: hi-low lines and up-down bars with nothing styled, so Office shows its defaults.
        group = (f'<c:stockChart>{series_xml("stock", case, flavour)}{dl}<c:hiLowLines/>'
                 f'<c:upDownBars><c:gapWidth val="150"/><c:upBars/><c:downBars/></c:upDownBars>'
                 f'<c:axId val="10"/><c:axId val="20"/></c:stockChart>')
        axes = axis_xml("c:catAx", 10, 20, "b", case, flavour, "catAx") + axis_xml("c:valAx", 20, 10, "l", case, flavour, "valAx")
    elif kind == "radar":
        group = (f'<c:radarChart><c:radarStyle val="{case.get("radarStyle", "marker")}"/>{vary}{series_xml("radar", case, flavour)}{dl}'
                 f'<c:axId val="10"/><c:axId val="20"/></c:radarChart>')
        axes = axis_xml("c:catAx", 10, 20, "b", case, flavour, "catAx") + axis_xml("c:valAx", 20, 10, "l", case, flavour, "valAx")
    elif kind == "combo":
        bars = case["series"][: case["split"]]
        lines = case["series"][case["split"]:]
        group = (f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>{series_xml("bar", case, flavour, 0, bars)}'
                 f'<c:gapWidth val="219"/><c:overlap val="-27"/><c:axId val="10"/><c:axId val="20"/></c:barChart>')
        sec = case.get("secondary", True)
        lid = (30, 40) if sec else (10, 20)
        group += (f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{series_xml("line", case, flavour, len(bars), lines)}'
                  f'<c:marker val="1"/><c:axId val="{lid[0]}"/><c:axId val="{lid[1]}"/></c:lineChart>')
        axes = axis_xml("c:catAx", 10, 20, "b", case, flavour, "catAx") + axis_xml("c:valAx", 20, 10, "l", case, flavour, "valAx")
        if sec:
            axes += axis_xml("c:catAx", 30, 40, "b", case, flavour, "cat2", deleted=True) + axis_xml(
                "c:valAx", 40, 30, "r", dict(case, val2=dict(case.get("val2", {}), major=False, crosses="max")), flavour, "val2")
    title = ""
    if case.get("title") is not None:
        tsz = case.get("titleSz", 1862 if flavour == "office" else None)
        t = case["title"]
        if t == "auto":
            title = '<c:title><c:overlay val="0"/>' + txpr(tsz, flavour=flavour) + '</c:title><c:autoTitleDeleted val="0"/>'
        else:
            rpr = f'<a:defRPr sz="{tsz}" b="0">{GRAY_TX if flavour == "office" else ""}</a:defRPr>' if tsz else '<a:defRPr/>'
            title = (f'<c:title><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{rpr}</a:pPr><a:r><a:rPr lang="en-US"/><a:t>{escape(t)}</a:t></a:r></a:p></c:rich></c:tx>'
                     f'<c:overlay val="0"/></c:title><c:autoTitleDeleted val="0"/>')
    else:
        title = '<c:autoTitleDeleted val="1"/>'
    legend = ""
    if case.get("legend", "b" if flavour == "office" else "r"):
        lp = case.get("legend", "b" if flavour == "office" else "r")
        legend = f'<c:legend><c:legendPos val="{lp}"/><c:overlay val="0"/>' + (
            '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln><a:effectLst/></c:spPr>' if flavour == "office" else "") + txpr(
            case.get("legendSz", 1197 if flavour == "office" else None), flavour=flavour) + '</c:legend>'
    plot_sppr = '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln><a:effectLst/></c:spPr>' if flavour == "office" else ""
    if flavour == "office":
        cs_sppr = case.get("chartSpPr", '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln><a:effectLst/></c:spPr>')
        cs_tx = '<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr/></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>'
    else:
        cs_sppr = case.get("chartSpPr", "")
        cs_tx = ""
    style = f'<c:style val="{case["style"]}"/>' if case.get("style") else ""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<c:chartSpace {NS}><c:date1904 val="0"/><c:lang val="en-US"/>'
            f'<c:roundedCorners val="0"/>{style}<c:chart>{title}<c:plotArea>{plot_layout}{group}{axes}{plot_sppr}</c:plotArea>{legend}'
            f'<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>{cs_sppr}{cs_tx}</c:chartSpace>')


def s(name, values, **kw):
    return dict(name=name, values=values, **kw)


Q = ["Q1", "Q2", "Q3", "Q4"]
M = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
S3 = [s("Series 1", [4.3, 2.5, 3.5, 4.5]), s("Series 2", [2.4, 4.4, 1.8, 2.8]), s("Series 3", [2, 2, 3, 5])]


def cases():
    C = []
    add = lambda **c: C.append(c)
    # --- columns / bars: PowerPoint's default inserted chart and variations
    add(id="col-default", kind="bar", cats=["Category 1", "Category 2", "Category 3", "Category 4"], series=S3, title="auto")
    add(id="col-notitle", kind="bar", cats=Q, series=S3)
    add(id="col-stacked", kind="bar", grouping="stacked", cats=Q, series=S3, title="Stacked")
    add(id="col-percent", kind="bar", grouping="percentStacked", cats=Q, series=S3, valAx={"fmt": "0%"})
    add(id="bar-clustered", kind="bar", dir="bar", cats=Q, series=S3, title="Horizontal", legend="r")
    add(id="bar-stacked-neg", kind="bar", dir="bar", grouping="stacked", cats=Q,
        series=[s("Gain", [5, 3, 6, 2]), s("Loss", [-2, -4, -1, -3])])
    add(id="col-negative", kind="bar", cats=M[:6], series=[s("Net", [12, -7, 3.5, -15, 22, 8])], legend=None)
    add(id="col-thousands", kind="bar", cats=Q, series=[s("Revenue", [125000, 180500, 97250, 210000]), s("Cost", [80000, 95000, 101000, 130500])],
        valAx={"fmt": "#,##0"}, title="Revenue vs cost")
    add(id="col-currency-dlbl", kind="bar", cats=Q, series=[s("Sales", [1250.5, 980.25, 1730, 1410])], legend=None,
        dLbls={"pos": "outEnd", "fmt": "$#,##0"}, valAx={"fmt": "$#,##0"})
    add(id="col-dlbl-ctr-stacked", kind="bar", grouping="stacked", cats=Q, series=S3, dLbls={"pos": "ctr"})
    add(id="col-gap50-ovl", kind="bar", cats=Q, series=S3, gap=50, overlap=30)
    add(id="col-gap500", kind="bar", cats=Q, series=S3[:1], gap=500, legend=None)
    add(id="col-small-values", kind="bar", cats=Q, series=[s("Rate", [0.012, 0.034, 0.027, 0.041])], legend=None)
    add(id="col-big-range", kind="bar", cats=Q, series=[s("A", [97, 98, 99, 100])], legend=None)
    add(id="col-fixed-scale", kind="bar", cats=Q, series=S3, valAx={"min": 0, "max": 10, "majorUnit": 2.5})
    add(id="col-log", kind="bar", cats=Q, series=[s("Hits", [3, 45, 900, 12000])], valAx={"logBase": 10}, legend=None)
    add(id="col-reversed", kind="bar", cats=Q, series=S3, catAx={"orientation": "maxMin"}, valAx={"orientation": "maxMin"})
    add(id="col-long-labels", kind="bar", cats=["North America region", "Europe, Middle East and Africa", "Asia Pacific and Japan", "Latin America"],
        series=S3[:2])
    add(id="col-many-cats", kind="bar", cats=[f"Item {i+1}" for i in range(30)], series=[s("V", [((i * 37) % 23) + 2 for i in range(30)])], legend=None)
    add(id="col-dates", kind="bar", cats=["2024-01", "2024-02", "2024-03", "2024-04", "2024-05", "2024-06", "2024-07", "2024-08"],
        series=[s("Units", [5, 8, 6, 9, 12, 10, 14, 13])], legend=None)
    add(id="col-axis-titles", kind="bar", cats=Q, series=S3, title="With axis titles", catAx={"title": "Quarter"}, valAx={"title": "Units (k)"}, legend="r")
    add(id="col-legend-top", kind="bar", cats=Q, series=S3, legend="t")
    add(id="col-legend-left", kind="bar", cats=Q, series=S3, legend="l")
    add(id="col-legend-tr", kind="bar", cats=Q, series=S3, legend="tr")
    add(id="col-small-frame", kind="bar", cats=Q, series=S3, title="Small", size=(3.5, 2.4))
    add(id="col-tall-frame", kind="bar", cats=Q, series=S3, size=(4, 6.5))
    add(id="col-manual-plot", kind="bar", cats=Q, series=S3, plotLayout=(0.2, 0.1, 0.6, 0.6))
    add(id="col-manual-outer", kind="bar", cats=Q, series=S3, plotLayout=(0.1, 0.05, 0.85, 0.75), layoutTarget="outer")
    add(id="col-ticks-cross", kind="bar", cats=Q, series=S3, catAx={"majorTick": "out"}, valAx={"majorTick": "cross", "line": True, "minor": True})
    add(id="col-arabic", kind="bar", cats=["الربع الأول", "الربع الثاني", "الربع الثالث", "الربع الرابع"],
        series=[s("المبيعات", [4.3, 2.5, 3.5, 4.5]), s("التكاليف", [2.4, 4.4, 1.8, 2.8])], title="المبيعات حسب الربع", script="arabic")
    add(id="col-many-series", kind="bar", cats=Q, series=[s(f"S{i+1}", [(i * 3 + j * 5) % 9 + 1 for j in range(4)]) for i in range(9)], legend="r")
    # Legend cap and wrap: what Office does with 20 series and with long names (the engine's third-of-the-chart cap is an assumption until measured).
    add(id="col-20-series-legend-r", kind="bar", cats=Q, series=[s(f"Series {i+1}", [(i * 3 + j * 5) % 9 + 1 for j in range(4)]) for i in range(20)], legend="r")
    add(id="col-20-series-legend-b", kind="bar", cats=Q, series=[s(f"Series {i+1}", [(i * 3 + j * 5) % 9 + 1 for j in range(4)]) for i in range(20)], legend="b")
    add(id="col-long-names-legend-r", kind="bar", cats=Q, series=[s(f"Quarterly revenue from the region number {i+1} excluding returns", [1 + i, 2 + i, 3 + i, 4 + i]) for i in range(3)], legend="r")
    add(id="col-bare", kind="bar", flavour="bare", cats=Q, series=S3)
    add(id="col-bare-style", kind="bar", flavour="bare", cats=Q, series=S3, style=26, title="Style 26")
    add(id="bar-bare", kind="bar", dir="bar", flavour="bare", cats=Q, series=S3)
    add(id="col-bare-dlbl", kind="bar", flavour="bare", cats=Q, series=S3[:1], dLbls={"val": True})
    add(id="col-chartarea-fill", kind="bar", cats=Q, series=S3,
        chartSpPr='<c:spPr><a:solidFill><a:srgbClr val="F2F2F2"/></a:solidFill><a:ln w="12700"><a:solidFill><a:srgbClr val="404040"/></a:solidFill></a:ln></c:spPr>')
    add(id="col-text-10", kind="bar", cats=Q, series=S3, title="10pt text", titleSz=1400,
        catAx={"sz": 1000}, valAx={"sz": 1000}, legendSz=1000)
    # --- lines / areas
    add(id="line-default", kind="line", cats=["Category 1", "Category 2", "Category 3", "Category 4"], series=S3, title="auto")
    add(id="line-markers", kind="line", cats=Q, series=S3, marker=("circle", 5))
    add(id="line-months", kind="line", cats=M, series=[s("2023", [3, 4, 6, 9, 12, 15, 17, 16, 13, 9, 6, 4]), s("2024", [4, 5, 7, 10, 14, 17, 19, 18, 14, 10, 7, 5])])
    add(id="line-smooth", kind="line", cats=M[:6], series=[s("Smooth", [1, 5, 2, 6, 3, 7])], smooth=True, legend=None)
    add(id="line-stacked", kind="line", grouping="stacked", cats=Q, series=S3, marker=("square", 6))
    add(id="line-gaps", kind="line", cats=M[:6], series=[s("Gappy", [3, None, 5, 6, None, 4])], legend=None, marker=("diamond", 7))
    add(id="line-dlbl-top", kind="line", cats=Q, series=S3[:1], dLbls={"pos": "t"}, marker=("circle", 5), legend=None)
    add(id="line-bare", kind="line", flavour="bare", cats=Q, series=S3)
    add(id="area-standard", kind="area", cats=Q, series=S3)
    add(id="area-stacked", kind="area", grouping="stacked", cats=M[:6], series=[s("A", [1, 2, 3, 2, 4, 3]), s("B", [2, 2, 1, 3, 2, 4])])
    add(id="area-percent", kind="area", grouping="percentStacked", cats=Q, series=S3)
    # --- pies / doughnuts
    pie_cats = ["1st Qtr", "2nd Qtr", "3rd Qtr", "4th Qtr"]
    add(id="pie-default", kind="pie", cats=pie_cats, series=[s("Sales", [8.2, 3.2, 1.4, 1.2])], title="auto")
    add(id="pie-labels-bestfit", kind="pie", cats=pie_cats, series=[s("Sales", [8.2, 3.2, 1.4, 1.2])], dLbls={"pos": "bestFit", "pct": True, "val": False, "cat": True}, legend=None)
    add(id="pie-outend", kind="pie", cats=pie_cats, series=[s("Sales", [8.2, 3.2, 1.4, 1.2])], dLbls={"pos": "outEnd", "val": True}, legend="r")
    add(id="pie-explode-rot", kind="pie", cats=pie_cats, series=[s("Sales", [8.2, 3.2, 1.4, 1.2])], explode={1: 15}, firstSliceAng=90)
    add(id="pie-many", kind="pie", cats=[f"P{i+1}" for i in range(9)], series=[s("Share", [30, 20, 12, 10, 8, 7, 6, 4, 3])], legend="r")
    # Busy pie: where Office puts ten labels, and whether it draws leader lines (dLbls() writes showLeaderLines for pies).
    add(id="pie-busy-outend", kind="pie", cats=[f"Category {i+1}" for i in range(10)], series=[s("Share", [40, 20, 10, 6, 5, 4, 4, 4, 4, 3])],
        dLbls={"pos": "outEnd", "pct": True, "val": False, "cat": True}, legend=None)
    add(id="pie-bare", kind="pie", flavour="bare", cats=pie_cats, series=[s("Sales", [8.2, 3.2, 1.4, 1.2])])
    add(id="doughnut-default", kind="doughnut", cats=pie_cats, series=[s("Sales", [8.2, 3.2, 1.4, 1.2])], hole=75, title="auto")
    add(id="doughnut-hole50-2ser", kind="doughnut", cats=pie_cats, series=[s("2023", [8.2, 3.2, 1.4, 1.2]), s("2024", [5, 4, 3, 2])], hole=50)
    # --- xy
    add(id="scatter-markers", kind="scatter", scatterStyle="lineMarker", cats=["", "", "", "", ""], xs=[0.7, 1.8, 2.6, 3.9, 5.2],
        series=[s("Y", [2.7, 3.2, 0.8, 4.1, 2.2])], marker=("circle", 5), lineW=None, title="auto")
    add(id="scatter-lines", kind="scatter", scatterStyle="lineMarker", cats=["", "", "", "", "", ""], xs=[0, 10, 20, 35, 50, 80],
        series=[s("A", [1, 4, 9, 16, 25, 36]), s("B", [3, 3, 8, 12, 30, 28])], marker=("square", 5))
    add(id="scatter-smooth", kind="scatter", scatterStyle="smoothMarker", smooth=True, cats=[""] * 6, xs=[-3, -1, 0, 1, 2, 4],
        series=[s("f", [9, 1, 0, 1, 4, 16])], legend=None)
    add(id="bubble", kind="bubble", cats=[""] * 4, xs=[1, 2, 3, 4.5], series=[s("B", [10, 25, 15, 30], sizes=[4, 9, 2, 12])], legend=None)
    add(id="radar", kind="radar", cats=["Speed", "Power", "Range", "Cost", "Comfort"], series=[s("A", [4, 3, 5, 2, 4]), s("B", [3, 5, 2, 4, 3])], marker=("circle", 5))
    add(id="radar-filled", kind="radar", radarStyle="filled", cats=["Speed", "Power", "Range", "Cost", "Comfort"], series=[s("A", [4, 3, 5, 2, 4])], legend=None)
    # --- combos
    add(id="combo-bar-line-2ax", kind="combo", split=2, cats=Q, series=[s("Units", [120, 150, 90, 170]), s("Returns", [20, 35, 15, 25]), s("Margin %", [0.21, 0.24, 0.18, 0.27])],
        val2={"fmt": "0%"}, title="Units and margin")
    add(id="combo-bar-line-1ax", kind="combo", split=1, secondary=False, cats=M[:6], series=[s("Actual", [5, 7, 6, 9, 8, 10]), s("Target", [6, 6, 7, 7, 8, 8])])
    # --- Arabic: `script="arabic"` marks the cases scored apart from the Latin ones (they measure joined-letter
    # widths, the complex-script face and right-to-left text). Appended last so the existing cases keep their slides.
    AR_Q = ["الربع الأول", "الربع الثاني", "الربع الثالث", "الربع الرابع"]
    add(id="ar-col-clustered", kind="bar", script="arabic", cats=AR_Q, series=S3[:2], legend="b")
    add(id="ar-legend-right", kind="bar", script="arabic", cats=Q, legend="r",
        series=[s("المبيعات", [4.3, 2.5, 3.5, 4.5]), s("التكاليف الإجمالية", [2.4, 4.4, 1.8, 2.8]), s("صافي الربح", [2, 2, 3, 5])])
    add(id="ar-pie-labels", kind="pie", script="arabic", cats=["الشرق الأوسط", "أوروبا", "آسيا والمحيط الهادئ", "أمريكا الشمالية"],
        series=[s("الحصة", [8.2, 3.2, 1.4, 1.2])], dLbls={"pos": "bestFit", "pct": True, "val": False, "cat": True}, legend="r")
    add(id="ar-title-wrap", kind="bar", script="arabic", cats=Q, series=S3[:1], size=(4.5, 3.2), legend=None,
        title="تقرير المبيعات الفصلية للمنطقة الشرقية مقارنة بالعام الماضي")
    add(id="ar-digits-cats", kind="bar", script="arabic", cats=["الربع 1 2024", "الربع ٢ 2024", "Q3 الثالث", "الشهر 12"], series=S3[:1], legend=None)
    add(id="ar-col-reversed", kind="bar", script="arabic", cats=AR_Q, series=S3[:2], catAx={"orientation": "maxMin"})
    # --- Guesses: rules the engine assumed without an Office case (DECISION_LOG §73), one or more cases each.
    # `guess` names the rule; appended last so the earlier cases keep their slides.
    names6 = ["Northern region sales", "Southern region sales", "Eastern region sales", "Western region sales", "Central region sales", "Overseas sales"]
    add(id="g-legend-r-mid", guess="side legend cap", kind="bar", cats=Q, legend="r",
        series=[s(n, [(i * 3 + j * 5) % 9 + 1 for j in range(4)]) for i, n in enumerate(names6)])
    add(id="g-legend-b-12", guess="legend grid", kind="bar", cats=Q, legend="b",
        series=[s(f"Region {i+1}", [(i * 3 + j * 5) % 9 + 1 for j in range(4)]) for i in range(12)])
    add(id="g-legend-t-long", guess="legend grid", kind="bar", cats=Q, legend="t",
        series=[s(f"Product line number {i+1}", [(i * 2 + j * 3) % 7 + 1 for j in range(4)]) for i in range(8)])
    add(id="g-pie-bestfit-long", guess="pie inside wrap", kind="pie", legend=None,
        cats=["Northern Europe and Scandinavia", "Southern Europe", "North Africa and Levant", "Gulf states"],
        series=[s("Share", [8.2, 3.2, 1.4, 1.2])], dLbls={"pos": "bestFit", "pct": True, "val": False, "cat": True})
    # Date axes: 45292 is 2024-01-01. Office picks the base unit itself (no baseTimeUnit written).
    day = lambda serials: [str(v) for v in serials]
    days10 = list(range(45292, 45302))
    add(id="g-date-days", guess="date axis", kind="bar", cats=day(days10), catNums=days10, catAx={"tag": "c:dateAx"},
        series=[s("Units", [5, 8, 6, 9, 12, 10, 14, 13, 11, 15])], legend=None)
    gaps = [45292, 45293, 45323, 45324]
    add(id="g-date-gaps", guess="date axis", kind="line", cats=day(gaps), catNums=gaps, catAx={"tag": "c:dateAx"},
        series=[s("Units", [5, 8, 6, 9])], legend=None, marker=("circle", 5))
    every28 = [45292 + 28 * k for k in range(6)]
    add(id="g-date-28", guess="date axis", kind="bar", cats=day(every28), catNums=every28, catAx={"tag": "c:dateAx"},
        series=[s("Units", [5, 8, 6, 9, 7, 10])], legend=None)
    months = [45292, 45323, 45352, 45383, 45413, 45444, 45474, 45505, 45536, 45566, 45597, 45627]
    add(id="g-date-months", guess="date axis", kind="line", cats=day(months), catNums=months, catAx={"tag": "c:dateAx"},
        series=[s("Units", [3, 4, 6, 9, 12, 15, 17, 16, 13, 9, 6, 4])], legend=None)
    years = [42736, 43101, 43466, 43831, 44197, 44562, 44927, 45292]
    add(id="g-date-years", guess="date axis", kind="bar", cats=day(years), catNums=years, catAx={"tag": "c:dateAx"},
        series=[s("Units", [5, 8, 6, 9, 12, 10, 14, 13])], legend=None)
    add(id="g-date-general", guess="date axis", kind="bar", cats=day(days10[:6]), catNums=days10[:6], catFmt="General",
        catAx={"tag": "c:dateAx"}, series=[s("Units", [5, 8, 6, 9, 12, 10])], legend=None)
    # Grouped categories: years over quarters.
    yq = "".join(f'<c:pt idx="{i}"><c:v>Q{i % 4 + 1}</c:v></c:pt>' for i in range(8))
    multi = ('<c:cat><c:multiLvlStrRef><c:f>Sheet1!$A$2:$B$9</c:f><c:multiLvlStrCache><c:ptCount val="8"/>'
             f'<c:lvl>{yq}</c:lvl><c:lvl><c:pt idx="0"><c:v>2023</c:v></c:pt><c:pt idx="4"><c:v>2024</c:v></c:pt></c:lvl>'
             '</c:multiLvlStrCache></c:multiLvlStrRef></c:cat>')
    add(id="g-multilevel", guess="grouped categories", kind="bar", cats=[f"Q{i % 4 + 1}" for i in range(8)], catXml=multi,
        series=[s("Units", [5, 8, 6, 9, 7, 10, 8, 12])], legend=None)
    # Trendlines: the bare one shows Office's default look; the labelled ones its equation text.
    xs6 = [1, 2, 3, 4, 5, 6]
    add(id="g-trend-linear-bare", guess="trendline look", kind="scatter", flavour="bare", cats=[""] * 6, xs=xs6,
        series=[s("Y", [2.1, 3.9, 6.2, 7.8, 10.1, 12.2])], marker=("circle", 5), scatterStyle="marker",
        serExtra='<c:trendline><c:trendlineType val="linear"/><c:dispRSqr val="0"/><c:dispEq val="0"/></c:trendline>', legend=None)
    add(id="g-trend-poly-eq", guess="trendline label", kind="scatter", cats=[""] * 6, xs=xs6, scatterStyle="marker", lineW=None,
        series=[s("Y", [9, 4, 1, 1, 4, 9])], marker=("circle", 5), legend=None,
        serExtra='<c:trendline><c:trendlineType val="poly"/><c:order val="2"/><c:dispRSqr val="1"/><c:dispEq val="1"/><c:trendlineLbl><c:numFmt formatCode="General" sourceLinked="0"/></c:trendlineLbl></c:trendline>')
    add(id="g-trend-movavg", guess="trendline look", kind="line", flavour="bare", cats=M[:8],
        series=[s("V", [3, 7, 4, 8, 5, 9, 6, 10])],
        serExtra='<c:trendline><c:trendlineType val="movingAvg"/><c:period val="2"/><c:dispRSqr val="0"/><c:dispEq val="0"/></c:trendline>')
    # Error bars, unstyled: their length rule and Office's default look.
    err = lambda kind, val: f'<c:errBars><c:errBarType val="both"/><c:errValType val="{kind}"/><c:noEndCap val="0"/><c:val val="{val}"/></c:errBars>'
    add(id="g-err-stddev", guess="error bars", kind="bar", flavour="bare", cats=Q, series=[s("V", [2, 4, 6, 8])], serExtra=err("stdDev", 1))
    add(id="g-err-stderr", guess="error bars", kind="bar", flavour="bare", cats=Q, series=[s("V", [2, 4, 6, 8])],
        serExtra='<c:errBars><c:errBarType val="both"/><c:errValType val="stdErr"/><c:noEndCap val="0"/></c:errBars>')
    add(id="g-err-fixed-pct", guess="error bars", kind="bar", flavour="bare", cats=Q,
        series=[s("A", [2, 4, 6, 8]), s("B", [3, 5, 7, 9])], serExtra=[err("fixedVal", 1), err("percentage", 10)])
    # Stock chart with nothing styled: Office's own hi-low line and up-down bar defaults.
    add(id="g-stock-bare", guess="stock defaults", kind="stock", flavour="bare", cats=["D1", "D2", "D3", "D4", "D5", "D6"], legend=None,
        series=[s("Open", [20, 23, 21.5, 22, 25, 23]), s("High", [24, 25.5, 22.5, 26, 27, 24]),
                s("Low", [18.5, 21, 19, 21.5, 22.5, 20]), s("Close", [23, 21.5, 22, 25, 23, 20.5])])
    # Value from cells, with the value: which comes first in the label.
    cells = ["Low", "Mid", "High", "Peak"]
    cache = "".join(f'<c:pt idx="{i}"><c:v>{t}</c:v></c:pt>' for i, t in enumerate(cells))
    add(id="g-dlbl-cells", guess="value from cells", kind="bar", cats=Q, series=[s("V", [2, 4, 6, 8])], legend=None,
        serExtra=('<c:dLbls><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/>'
                  '<c:showPercent val="0"/><c:showBubbleSize val="0"/><c:extLst><c:ext uri="{CE6537A1-D6FC-4f65-9D91-7224C49458BB}" '
                  'xmlns:c15="http://schemas.microsoft.com/office/drawing/2012/chart"><c15:showDataLabelsRange val="1"/></c:ext></c:extLst></c:dLbls>'),
        serExt=('<c:extLst><c:ext uri="{02D57815-91ED-43cb-92C2-25804820EDAC}" xmlns:c15="http://schemas.microsoft.com/office/drawing/2012/chart">'
                f'<c15:datalabelsRange><c15:f>Sheet1!$D$2:$D$5</c15:f><c15:dlblRangeCache><c:ptCount val="4"/>{cache}</c15:dlblRangeCache></c15:datalabelsRange></c:ext></c:extLst>'))
    # A pie whose labels were dragged off their slices: Office's own leader lines (unstyled).
    moved = "".join(f'<c:dLbl><c:idx val="{i}"/><c:layout><c:manualLayout><c:x val="{dx}"/><c:y val="{dy}"/></c:manualLayout></c:layout>'
                    '<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>'
                    for i, (dx, dy) in enumerate([(0.08, -0.08), (0.08, 0.08), (-0.08, 0.06), (-0.08, -0.06)]))
    add(id="g-pie-leaders", guess="leader lines", kind="pie", flavour="bare", cats=pie_cats, series=[s("Sales", [8.2, 3.2, 1.4, 1.2])], legend=None,
        serExtra=(f'<c:dLbls>{moved}<c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/>'
                  '<c:showPercent val="0"/><c:showBubbleSize val="0"/><c:showLeaderLines val="1"/></c:dLbls>'))
    return C


def build_pptx(cases, path, slide_w=13.333, slide_h=7.5):
    prs = Presentation()
    prs.slide_width = Inches(slide_w)
    prs.slide_height = Inches(slide_h)
    blank = prs.slide_layouts[6]
    frames = []
    for c in cases:
        slide = prs.slides.add_slide(blank)
        w, h = c.get("size", (8.0, 4.5))
        x, y = 0.75, 0.75
        cd = CategoryChartData()
        cd.categories = ["a", "b"]
        cd.add_series("s", (1, 2))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        part = gf.chart.part
        xml = chart_xml(c)
        # Keep python-pptx's embedded-workbook relationship id for externalData.
        part._element = etree.fromstring(xml.encode("utf-8"))
        # Every chart gets a caption below the frame, outside it, naming the case.
        tb = slide.shapes.add_textbox(Inches(0.75), Inches(slide_h - 0.6), Inches(8), Inches(0.4))
        tb.text_frame.text = c["id"]
        frames.append(dict(id=c["id"], x=x, y=y, w=w, h=h))
    prs.save(path)
    return frames


def write_case_data(ws, case):
    """Put the case's data in the cells `series_xml` references, so Excel, which
    recalculates chart links on open, draws the same numbers as the cache."""
    xy = case["kind"] in ("scatter", "bubble")
    for r, v in enumerate(case["xs"] if xy else case["cats"]):
        if v not in (None, ""):
            ws.cell(row=r + 2, column=1, value=v)
    for i, s in enumerate(case["series"]):
        ws.cell(row=1, column=i + 2, value=s["name"])
        for r, v in enumerate(s["values"]):
            if v is not None:
                ws.cell(row=r + 2, column=i + 2, value=v)
        for r, v in enumerate(s.get("sizes", [])):
            ws.cell(row=r + 2, column=26, value=v)


def build_xlsx(cases, path):
    """One sheet per case; openpyxl draws a placeholder bar chart we then swap for the case XML."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    wb = Workbook()
    wb.remove(wb.active)
    for i, c in enumerate(cases):
        ws = wb.create_sheet(f"c{i+1}")
        write_case_data(ws, c)
        # The case id sits clear of the data block (A..Z).
        ws["AB1"] = c["id"]
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=2, max_row=3))
        w, h = c.get("size", (8.0, 4.5))
        ch.width = w * 2.54
        ch.height = h * 2.54
        ws.add_chart(ch, "B3")
    tmp = path + ".tmp"
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            name = item.filename
            if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                n = int(name[len("xl/charts/chart"):-4])
                c = cases[n - 1]
                # Quoted: a bare `c14!` would read as a cell reference.
                cx = dict(c, sheetRef=f"'c{n}'")
                # Excel's own inserted-chart defaults: 9pt text, 14pt title, white chart area.
                if cx.get("flavour", "office") == "office":
                    cx.setdefault("titleSz", 1400)
                    for ax in ("catAx", "valAx"):
                        cx[ax] = dict(cx.get(ax, {}))
                        cx[ax].setdefault("sz", 900)
                    cx.setdefault("legendSz", 900)
                    if cx.get("dLbls"):
                        cx["dLbls"] = dict(cx["dLbls"], sz=cx["dLbls"].get("sz", 900))
                    cx.setdefault("chartSpPr", '<c:spPr><a:solidFill><a:schemeClr val="bg1"/></a:solidFill>' + GRID_LN + '<a:effectLst/></c:spPr>')
                data = chart_xml(cx).encode("utf-8")
            zout.writestr(item, data)
    os.remove(tmp)


if __name__ == "__main__":
    # Next to this script by default: export-office.ps1, render.ts and score.py read from here.
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
    os.makedirs(out, exist_ok=True)
    cs = cases()
    frames = build_pptx(cs, os.path.join(out, "chart-oracle.pptx"))
    build_xlsx(cs, os.path.join(out, "chart-oracle.xlsx"))
    with open(os.path.join(out, "cases.json"), "w", encoding="utf-8") as f:
        json.dump(dict(cases=cs, frames=frames), f, indent=1, ensure_ascii=False)
    print(len(cs), "cases")
